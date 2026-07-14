#!/usr/bin/env python3
"""
Aravest Presentation Design System — Calibration Set v3
Content verbatim/near-verbatim from the LOCKED Board Strategy Narrative v8.

v2 changes vs v1:
- Gold demoted to a rare signal: one element max per slide; working slides carry none.
- Secondary display face: Spectral (website-aligned, OFL) for covers + one headline
  phrase only. Core face remains Aptos Display / Aptos.
- Flat geometry: no shadows, no card grids, no accent strips; hairline rules,
  controlled tonal fills, typography-led hierarchy.
- Two cover directions (A: full-bleed cinematic + navy overlay; B: navy + vertical crop).
- Three design rhythms: editorial (covers, exec summary), analysis (pathways, AUM,
  pipeline), governance (board decision).
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

A = "assets"

# ---------- approved reconciled palette ----------
NAVY      = RGBColor(0x00,0x2B,0x5C)
DEEP_NAVY = RGBColor(0x00,0x21,0x47)
GRAPHITE  = RGBColor(0x35,0x46,0x4F)
SLATE     = RGBColor(0x86,0x93,0x97)
STEEL     = RGBColor(0x46,0x82,0xB4)
TEAL      = RGBColor(0x00,0x80,0x80)
GOLD      = RGBColor(0xC8,0xA9,0x51)
MIST      = RGBColor(0xE8,0xE8,0xE8)
MIST_LT   = RGBColor(0xF4,0xF5,0xF6)
WHITE     = RGBColor(0xFF,0xFF,0xFF)
WHITE_D   = RGBColor(0xD5,0xDD,0xE6)
TINT_NAVY = RGBColor(0xC3,0xCE,0xDC)   # tonal navy fill for schematic marks
TONAL_LN  = RGBColor(0x14,0x3A,0x6B)   # tonal facet line on deep navy

HEAD = "Aptos Display"
BODY = "Aptos"
SERIF   = "Spectral Light"   # secondary display face (website-aligned)
SERIF_R = "Spectral"

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)

def rect(s,x,y,w,h,fill=None,line=None,line_w=None,shape=MSO_SHAPE.RECTANGLE,alpha=None):
    sp=s.shapes.add_shape(shape,x,y,w,h); sp.shadow.inherit=False
    if fill is None: sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb=fill
        if alpha is not None:
            clr=sp.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            a=clr.makeelement(qn('a:alpha'),{'val':str(int(alpha*100000))}); clr.append(a)
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=line_w or Pt(1)
    return sp

def hrule(s,x,y,w,weight=0.75,color=MIST):
    return rect(s,x,y,w,Pt(weight),fill=color)

def _set(r,text,size,font,color,bold=False,italic=False,track=None,caps=False):
    r.text=text; r.font.size=Pt(size); r.font.name=font
    r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color
    rPr=r._r.get_or_add_rPr()
    for tag in ("a:latin","a:cs"):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",font)
    if track is not None: rPr.set("spc",str(int(track*100)))
    if caps: rPr.set("cap","all")

def textbox(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    return tb,tf

def para(tf,first=False):
    return tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()

def line(tf,text,size,font=BODY,color=GRAPHITE,bold=False,italic=False,align=PP_ALIGN.LEFT,
         sb=0,sa=6,first=False,track=None,caps=False,ls=None):
    p=para(tf,first=first); p.alignment=align
    p.space_before=Pt(sb); p.space_after=Pt(sa)
    if ls: p.line_spacing=ls
    r=p.add_run(); _set(r,text,size,font,color,bold=bold,italic=italic,track=track,caps=caps)
    return p

def furniture(s,page):
    tb,tf=textbox(s,Inches(0.7),Inches(7.06),Inches(5),Inches(0.3),anchor=MSO_ANCHOR.MIDDLE)
    line(tf,"Private and Confidential",8.5,color=SLATE,first=True,sa=0)
    tb,tf=textbox(s,Inches(11.6),Inches(7.06),Inches(1.03),Inches(0.3),anchor=MSO_ANCHOR.MIDDLE)
    line(tf,str(page),8.5,color=SLATE,align=PP_ALIGN.RIGHT,first=True,sa=0)
    s.shapes.add_picture(A+"/logo_navy.png",Inches(11.90),Inches(0.36),height=Inches(0.38))

def kicker(s,x,y,text,color=SLATE):
    tb,tf=textbox(s,x,y,Inches(9),Inches(0.28))
    line(tf,text,10,color=color,bold=True,first=True,track=1.6,caps=True,sa=0)

SUBTITLE=("From integration to institution-building: establishing the post-SMFL track record, "
          "defining the near-term investment focus and developing the pathways for sustainable "
          "long-term growth.")

def cover_title(s,x,y,w,white=True,size=44):
    """Editorial two-run title: roman + italic phrase (website treatment)."""
    col=WHITE if white else NAVY
    tb,tf=textbox(s,x,y,w,Inches(1.05))
    p=para(tf,first=True); p.alignment=PP_ALIGN.LEFT; p.space_after=Pt(0)
    r=p.add_run(); _set(r,"Aravest’s ",size,SERIF,col)
    r=p.add_run(); _set(r,"Next Phase",size,SERIF,col,italic=True)

# =====================================================================
# SLIDE 1 — COVER OPTION A · full-bleed cinematic, navy overlay, editorial
# =====================================================================
s=slide()
s.shapes.add_picture(A+"/coverA_full.jpg",0,0,width=SW,height=SH)
rect(s,0,0,SW,SH,fill=DEEP_NAVY,alpha=0.66)                    # controlled overlay
s.shapes.add_picture(A+"/logo_white.png",Inches(1.0),Inches(0.75),height=Inches(0.58))
hrule(s,Inches(1.0),Inches(2.98),Inches(0.85),weight=1.5,color=GOLD)   # the one gold element
tb,tf=textbox(s,Inches(1.0),Inches(3.16),Inches(10.5),Inches(0.32))
line(tf,"Board Strategy Meeting  ·  September 2026",11,color=WHITE_D,bold=True,
     first=True,track=1.8,caps=True,sa=0)
cover_title(s,Inches(0.98),Inches(3.56),Inches(11.4),white=True,size=46)
tb,tf=textbox(s,Inches(1.0),Inches(4.78),Inches(8.7),Inches(1.0))
line(tf,SUBTITLE,14.5,color=WHITE_D,first=True,ls=1.32,sa=0)
tb,tf=textbox(s,Inches(1.0),Inches(6.95),Inches(11.3),Inches(0.32))
line(tf,"Aravest Fund Management Pte. Ltd.      Private and Confidential",9,
     color=WHITE_D,first=True,sa=0)

# =====================================================================
# SLIDE 2 — COVER OPTION B · navy institutional, vertical crop, tonal facet
# =====================================================================
s=slide()
rect(s,0,0,SW,SH,fill=DEEP_NAVY)
# subtle tonal facet motif — one large outlined triangle, cropped low-left
rect(s,Inches(-2.1),Inches(4.15),Inches(6.4),Inches(4.6),fill=None,line=TONAL_LN,
     line_w=Pt(1.0),shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
# vertical architectural crop, right edge
s.shapes.add_picture(A+"/coverB_strip.jpg",Inches(9.36),0,width=Inches(3.973),height=SH)
s.shapes.add_picture(A+"/logo_white.png",Inches(1.0),Inches(0.75),height=Inches(0.58))
hrule(s,Inches(1.0),Inches(2.86),Inches(0.85),weight=1.5,color=GOLD)   # the one gold element
tb,tf=textbox(s,Inches(1.0),Inches(3.04),Inches(8),Inches(0.32))
line(tf,"Board Strategy Meeting  ·  September 2026",11,color=SLATE,bold=True,
     first=True,track=1.8,caps=True,sa=0)
cover_title(s,Inches(0.98),Inches(3.44),Inches(8.2),white=True,size=42)
tb,tf=textbox(s,Inches(1.0),Inches(4.58),Inches(7.5),Inches(1.1))
line(tf,SUBTITLE,13.5,color=WHITE_D,first=True,ls=1.32,sa=0)
tb,tf=textbox(s,Inches(1.0),Inches(6.95),Inches(8),Inches(0.32))
line(tf,"Aravest Fund Management Pte. Ltd.      Private and Confidential",9,
     color=WHITE_D,first=True,sa=0)

# =====================================================================
# SLIDE 3 — EXECUTIVE SUMMARY · editorial rhythm (no kicker, no gold)
# =====================================================================
s=slide()
rect(s,0,0,SW,SH,fill=WHITE)
LX=Inches(1.0); CW=Inches(11.33)
tb,tf=textbox(s,LX,Inches(0.60),CW,Inches(1.4))
p=para(tf,first=True); p.alignment=PP_ALIGN.LEFT; p.space_after=Pt(0); p.line_spacing=1.06
r=p.add_run(); _set(r,"Aravest is moving from integration",31,HEAD,NAVY)
p=tf.add_paragraph(); p.alignment=PP_ALIGN.LEFT; p.space_after=Pt(0); p.line_spacing=1.06
r=p.add_run(); _set(r,"into ",31,HEAD,NAVY)
r=p.add_run(); _set(r,"institution-building",31,SERIF,NAVY,italic=True)
r=p.add_run(); _set(r,".",31,HEAD,NAVY)
tb,tf=textbox(s,LX,Inches(1.98),Inches(11.0),Inches(0.85))
line(tf,"The initial integration into SMFL has been substantially completed. The next phase "
        "is to develop the institutional foundations and investment record required for "
        "sustainable long-term growth.",14,color=GRAPHITE,first=True,ls=1.28,sa=0)
SPINE_Y=Inches(3.16)
hrule(s,LX,SPINE_Y,CW,weight=1.0,color=MIST)
concl=[
 ("01","The Aravest Way",
  "Values-led in how it invests; value-focused in where it deploys.",
  "Investment management is a responsibility before it is a business — sector-flexible, but not strategy-neutral."),
 ("02","The post-SMFL record",
  "Judgement, execution, stewardship, performance, realisation and repeatability.",
  "Confidence moves from the individual opportunity, to the repeat relationship, to the Aravest strategy."),
 ("03","Two parallel pathways",
  "Diversified track-record building on the existing US$300 million.",
  "Institutional-scale opportunities, case-by-case, can accelerate AUM, relevance and visibility."),
 ("04","Institutional standards",
  "Committed equity bears performance; wider liabilities are ring-fenced.",
  "No automatic recourse to Aravest or SMFL; every fund structured to institutional standards."),
]
colw=Inches(2.66); gap=Inches(0.23)
for i,(no,h,b1,b2) in enumerate(concl):
    x=LX+i*(colw+gap)
    rect(s,x,SPINE_Y-Pt(2.4),Pt(1.6),Pt(6.5),fill=NAVY)
    tb,tf=textbox(s,x,SPINE_Y+Inches(0.18),colw,Inches(2.7))
    line(tf,no,10.5,color=SLATE,bold=True,first=True,track=1.2,sa=6)
    line(tf,h,15.5,font=HEAD,color=NAVY,sa=7,ls=1.03)
    line(tf,b1,12.5,color=GRAPHITE,ls=1.24,sa=7)
    line(tf,b2,12,color=SLATE,ls=1.24,sa=0)
rect(s,LX,Inches(6.28),RW if False else CW,Pt(1.5),fill=NAVY)
tb,tf=textbox(s,LX,Inches(6.42),CW,Inches(0.5),anchor=MSO_ANCHOR.MIDDLE)
p=para(tf,first=True); p.space_after=Pt(0); p.line_spacing=1.15
r=p.add_run(); _set(r,"One strategy.  ",13.5,HEAD,NAVY)
r=p.add_run(); _set(r,"The standards stay constant while the record, the pathways and the structure "
                     "compound into an institution.",13.5,BODY,GRAPHITE)
furniture(s,2)

# =====================================================================
# SLIDE 4 — TWO-PATHWAY STRATEGY · analysis rhythm, equal status, matrix
# =====================================================================
s=slide()
rect(s,0,0,SW,SH,fill=WHITE)
LX=Inches(0.7); RW=Inches(11.93)
kicker(s,LX,Inches(0.48),"Strategy · Two parallel pathways")
tb,tf=textbox(s,LX,Inches(0.80),Inches(11.9),Inches(0.6))
line(tf,"Aravest will build the platform through two complementary pathways.",
     24,font=HEAD,color=NAVY,first=True,ls=1.02,sa=0)

LABW=Inches(2.30); P1X=LX+LABW+Inches(0.25); COLW=Inches(4.55); P2X=P1X+COLW+Inches(0.35)
HDY=Inches(1.52)
for px,name,role in [(P1X,"Pathway 1 — Diversified Track Record Building","The principal programme"),
                     (P2X,"Pathway 2 — Institutional-Scale Opportunities","Case-by-case acceleration")]:
    rect(s,px,HDY,COLW,Pt(2),fill=NAVY)
    tb,tf=textbox(s,px,HDY+Inches(0.11),COLW,Inches(0.78))
    line(tf,name,15,font=HEAD,color=NAVY,first=True,ls=1.05,sa=3)
    line(tf,role.upper(),9.5,color=STEEL,bold=True,track=1.2,sa=0)
rows=[
 ("Primary purpose",
  "Build multiple proof points and a repeatable post-SMFL investment record.",
  "Accelerate AUM, institutional relevance and market visibility."),
 ("Typical opportunity",
  "Selected small and medium transactions across repeatable themes.",
  "Selected large transactions capable of creating meaningful scale."),
 ("Capital profile",
  "The existing US$300 million support, diversified across several investments.",
  "May require materially greater transaction-specific capital."),
 ("Risk profile",
  "More diversified, with lower single-asset concentration.",
  "Higher single-asset concentration and greater liquidity exposure."),
 ("Expected AUM effect",
  "Gradual and potentially uneven growth.",
  "Potential step-change growth."),
]
ry=Inches(2.54); rh=Inches(0.77)
for i,(lab,v1,v2) in enumerate(rows):
    if i>0: hrule(s,LX,ry,RW,weight=0.75,color=MIST)
    tb,tf=textbox(s,LX,ry+Inches(0.12),LABW,rh,anchor=MSO_ANCHOR.TOP)
    line(tf,lab.upper(),9.5,color=SLATE,bold=True,first=True,track=0.8,sa=0,ls=1.1)
    for px,v in [(P1X,v1),(P2X,v2)]:
        tb,tf=textbox(s,px,ry+Inches(0.10),COLW,rh)
        line(tf,v,13,color=GRAPHITE,first=True,ls=1.18,sa=0)
    ry+=rh
hrule(s,LX,ry,RW,weight=1.5,color=NAVY)
tb,tf=textbox(s,LX,ry+Inches(0.14),RW,Inches(0.45),anchor=MSO_ANCHOR.MIDDLE)
p=para(tf,first=True); p.space_after=Pt(0)
r=p.add_run(); _set(r,"Run in parallel:  ",13,BODY,SLATE)
r=p.add_run(); _set(r,"Pathway 1 builds the record; Pathway 2 can accelerate scale.",13,BODY,NAVY,bold=True)
furniture(s,3)

# =====================================================================
# SLIDE 5 — CAPITAL & AUM · analysis rhythm, analytical, no boxes, no gold
# =====================================================================
s=slide()
rect(s,0,0,SW,SH,fill=WHITE)
LX=Inches(0.7); RW=Inches(11.93)
kicker(s,LX,Inches(0.48),"Capital & AUM")
tb,tf=textbox(s,LX,Inches(0.80),Inches(11.9),Inches(0.6))
line(tf,"The two pathways are expected to produce different AUM profiles.",
     24,font=HEAD,color=NAVY,first=True,ls=1.02,sa=0)
hrule(s,LX,Inches(1.66),RW,weight=0.75,color=MIST)
tb,tf=textbox(s,LX,Inches(1.80),RW,Inches(0.44),anchor=MSO_ANCHOR.MIDDLE)
line(tf,"Closing AUM  =  Opening AUM  +  New investments  −  Realisations  ±  Valuation and FX movements",
     15,font=HEAD,color=NAVY,align=PP_ALIGN.CENTER,first=True,sa=0)
hrule(s,LX,Inches(2.38),RW,weight=0.75,color=MIST)

ANX=LX; ANW=Inches(3.05)
rect(s,ANX,Inches(2.78),Pt(2.5),Inches(2.35),fill=NAVY)
tb,tf=textbox(s,ANX+Inches(0.24),Inches(2.82),ANW-Inches(0.2),Inches(2.5))
line(tf,"SPONSOR CAPITAL SUPPORT",9,color=SLATE,bold=True,first=True,track=1.0,sa=7)
line(tf,"US$300m",34,font=HEAD,color=NAVY,sa=8,ls=1.0)
line(tf,"Existing support, diversified across several Pathway 1 investments.",12.5,
        color=GRAPHITE,ls=1.26,sa=6)
line(tf,"Warehousing bridges transaction execution and third-party capital formation.",12.5,
        color=GRAPHITE,ls=1.26,sa=0)

def schematic(x,w,title,heights,step_idx=None,accent=STEEL):
    base=Inches(4.86); maxh=1.30
    tb,tf=textbox(s,x,Inches(2.80),w,Inches(0.5))
    line(tf,title,13.5,font=HEAD,color=NAVY,first=True,ls=1.03,sa=0)
    n=len(heights); bw=Inches(0.40); gap=(w-bw*n)/(n-1)
    for i,hh in enumerate(heights):
        bh=Inches(hh*maxh)
        col=accent if (step_idx is not None and i>=step_idx) else TINT_NAVY
        rect(s,x+i*(bw+gap),base-bh,bw,bh,fill=col)
    hrule(s,x,base,w,weight=1.0,color=SLATE)
def schematic_caption(x,w,text):
    tb,tf=textbox(s,x,Inches(5.00),w,Inches(0.95))
    line(tf,text,11.5,color=GRAPHITE,first=True,ls=1.24,sa=4)
    line(tf,"Illustrative profile only — not a forecast.",9,color=SLATE,italic=True,sa=0)

S1X=Inches(4.20); S1W=Inches(3.85)
schematic(S1X,S1W,"Pathway 1 — gradual and uneven",[0.42,0.55,0.47,0.62,0.52,0.70,0.63])
schematic_caption(S1X,S1W,"AUM may grow gradually, remain flat where new investments replace "
                          "realisations, or decline temporarily after disposals or sell-downs.")
S2X=Inches(8.78); S2W=Inches(3.85)
schematic(S2X,S2W,"Pathway 2 — potential step-change",[0.36,0.38,0.40,0.93,0.95,1.0,1.0],step_idx=3)
schematic_caption(S2X,S2W,"A single institutional-scale transaction can create a material or "
                          "step-change increase in AUM and more meaningful ticket sizes.")

hrule(s,LX,Inches(6.18),RW,weight=0.75,color=MIST)
tb,tf=textbox(s,LX,Inches(6.30),RW,Inches(0.55))
p=para(tf,first=True); p.space_after=Pt(0); p.line_spacing=1.2
r=p.add_run(); _set(r,"PROOF POINT — CONRAD SEOUL   ",10,BODY,NAVY,bold=True,track=1.0)
r=p.add_run(); _set(r,"Sponsor-supported acquisition; institutional capital subsequently introduced "
                     "(GIC, M&G), recycling capital while retaining the management relationship.",12.5,BODY,GRAPHITE)
tb,tf=textbox(s,LX,Inches(6.86),RW,Inches(0.3))
line(tf,"Source: Locked Board Strategy Narrative v8. Financial figures to be populated by Finance "
        "before Board circulation.",8.5,color=SLATE,italic=True,first=True,sa=0)
furniture(s,4)

# =====================================================================
# SLIDE 6 — BOARD DECISION · governance rhythm, action-verb architecture
# =====================================================================
s=slide()
rect(s,0,0,SW,SH,fill=WHITE)
LX=Inches(0.7); RW=Inches(11.93)
rect(s,LX,Inches(0.53),Inches(0.11),Inches(0.11),fill=GOLD)   # the single gold element
tb,tf=textbox(s,LX+Inches(0.24),Inches(0.48),Inches(9),Inches(0.28))
line(tf,"For decision",10,color=NAVY,bold=True,first=True,track=1.6,caps=True,sa=0)
tb,tf=textbox(s,LX,Inches(0.82),Inches(12.0),Inches(1.0))
line(tf,"The Board is asked to endorse the two-pathway strategy and the basis on which "
        "capital is committed.",24,font=HEAD,color=NAVY,first=True,ls=1.05,sa=0)

decs=[
 ("Endorse","the two parallel pathways as the basis for building the Aravest platform — "
            "diversified track-record building alongside selective institutional-scale opportunities."),
 ("Support","deployment of the existing US$300 million across Pathway 1 to build multiple "
            "proof points and diversify risk."),
 ("Endorse","warehousing as an execution enabler, bridging the timing gap between securing "
            "a transaction and forming third-party capital."),
 ("Agree","that each Institutional-Scale Opportunity is brought to the Board separately — with its "
          "investment merits, capital requirement and risk profile — for decision case-by-case."),
]
DY=Inches(2.10); DH=Inches(0.86)
for i,(verb,rest) in enumerate(decs):
    if i>0: hrule(s,LX,DY,RW,weight=0.75,color=MIST)
    tb,tf=textbox(s,LX,DY+Inches(0.16),Inches(0.5),Inches(0.4))
    line(tf,f"{i+1:02d}",11,color=SLATE,bold=True,first=True,track=1.0,sa=0)
    tb,tf=textbox(s,LX+Inches(0.62),DY+Inches(0.13),Inches(1.62),Inches(0.5))
    line(tf,verb.upper(),15,font=HEAD,color=NAVY,first=True,track=1.0,sa=0)
    tb,tf=textbox(s,LX+Inches(2.40),DY+Inches(0.13),Inches(9.50),Inches(0.72))
    line(tf,rest,13.5,color=GRAPHITE,first=True,ls=1.2,sa=0)
    DY+=DH
GY=DY+Inches(0.14)
rect(s,LX,GY,RW,Pt(1.5),fill=NAVY)
tb,tf=textbox(s,LX,GY+Inches(0.13),RW,Inches(0.9))
line(tf,"GOVERNING PRINCIPLE",9,color=SLATE,bold=True,first=True,track=1.2,sa=5)
line(tf,"Committed equity is exposed to investment performance; wider fund and asset liabilities "
        "are generally ring-fenced, with no automatic recourse to Aravest as manager or SMFL as "
        "shareholder. Any exception must be explicit and separately approved.",13.5,color=NAVY,ls=1.24,sa=0)
tb,tf=textbox(s,LX,GY+Inches(1.20),RW,Inches(0.3))
line(tf,"Formal resolutions to be confirmed by management and the company secretary; this slide "
        "reflects the locked Board Strategy Narrative v8 and does not itself constitute a resolution.",
     8.5,color=SLATE,italic=True,first=True,sa=0)
furniture(s,5)

# =====================================================================
# SLIDE 7 — PIPELINE · analysis rhythm, rule-led institutional table
# =====================================================================
s=slide()
rect(s,0,0,SW,SH,fill=WHITE)
LX=Inches(0.7); RW=Inches(11.93)
kicker(s,LX,Inches(0.48),"Investment case · Current pipeline")
tb,tf=textbox(s,LX,Inches(0.80),Inches(11.9),Inches(0.6))
line(tf,"Current opportunities illustrate how the two pathways may develop in practice.",
     24,font=HEAD,color=NAVY,first=True,ls=1.02,sa=0)

C1=LX; C1W=Inches(3.15)
C2=Inches(4.15); C2W=Inches(4.35)
C3=Inches(8.80); C3W=Inches(3.83)
HY=Inches(1.74)
for x,w,h in [(C1,C1W,"Strategic role"),(C2,C2W,"Illustrative opportunities"),(C3,C3W,"Connection to the pathway")]:
    tb,tf=textbox(s,x,HY,w,Inches(0.3))
    line(tf,h.upper(),9.5,color=SLATE,bold=True,first=True,track=1.1,sa=0)
rect(s,LX,HY+Inches(0.34),RW,Pt(1.5),fill=NAVY)
prows=[
 ("Institutional-scale opportunity","PATHWAY 2",
  "Marina One and other selected large opportunities.",
  "Potential to create a material increase in AUM, institutional relevance and market visibility."),
 ("Portfolio building block","PATHWAY 1",
  "Project Guardian / PBSA and subsequent PBSA opportunities.",
  "Build a repeatable strategy through aggregation into a portfolio of institutional scale."),
 ("Track-record proof point","PATHWAY 1",
  "Courtyard by Marriott Suwon, ARA-NH Fund 2, Aberdeen / Shinyoung recapitalisation and "
  "other relevant transactions.",
  "Deepen post-SMFL execution evidence and create additional realisation or recycling pathways."),
]
ry=HY+Inches(0.46); rh=Inches(1.24)
for i,(role,tag,opp,conn) in enumerate(prows):
    if i>0: hrule(s,LX,ry,RW,weight=0.75,color=MIST)
    tb,tf=textbox(s,C1,ry+Inches(0.18),C1W,rh)
    line(tf,role,14,font=HEAD,color=NAVY,first=True,ls=1.06,sa=5)
    line(tf,tag,9,color=STEEL,bold=True,track=1.2,sa=0)
    tb,tf=textbox(s,C2,ry+Inches(0.20),C2W,rh)
    line(tf,opp,12,color=GRAPHITE,first=True,ls=1.24,sa=0)
    tb,tf=textbox(s,C3,ry+Inches(0.20),C3W,rh)
    line(tf,conn,12,color=GRAPHITE,first=True,ls=1.24,sa=0)
    ry+=rh
rect(s,LX,ry,RW,Pt(1.0),fill=NAVY)
tb,tf=textbox(s,LX,ry+Inches(0.14),RW,Inches(0.55))
line(tf,"Private credit is excluded as a core Aravest growth strategy (an agreed SMDAM strategy, "
        "with Aravest providing support where relevant). Transaction names, figures and status to "
        "be refreshed by the relevant transaction teams before Board circulation.",
     9,color=SLATE,italic=True,first=True,ls=1.24,sa=0)
furniture(s,6)

prs.save("Aravest_Calibration_Set_v3.pptx")
print("saved Aravest_Calibration_Set_v3.pptx,",len(prs.slides._sldIdLst),"slides")
