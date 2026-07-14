#!/usr/bin/env python3
"""
Aravest Design System — component library (V4-final direction, approved).

Final refinements applied:
- Pathway tints strengthened 5% -> 7% for projector readability; borders 35% -> 40%.
- Card body text minimum 13pt where possible.
- Cover B = default Board cover; Cover A = investor/marketing variant.
- Light decision = default; dark field = emphasis variant only.
- Dense tables use pathway markers, not full row tints.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
import json, os

A = "assets"

# ---------------- palette & functional roles ----------------
NAVY      = RGBColor(0x00,0x2B,0x5C)   # brand primary — headings, first data series
DEEP_NAVY = RGBColor(0x00,0x21,0x47)   # structural — dark fields, conclusions, principles
GRAPHITE  = RGBColor(0x35,0x46,0x4F)   # secondary structure & body text
SLATE     = RGBColor(0x86,0x93,0x97)   # quiet furniture: footnotes, page numbers
STEEL     = RGBColor(0x46,0x82,0xB4)   # Pathway 2 — institutional scale / acceleration
TEAL      = RGBColor(0x00,0x80,0x80)   # Pathway 1 — the record / repeatability
GOLD      = RGBColor(0xC8,0xA9,0x51)   # rare Board-level signal only
MIST      = RGBColor(0xE8,0xE8,0xE8)
WHITE     = RGBColor(0xFF,0xFF,0xFF)
WHITE_D   = RGBColor(0xD5,0xDD,0xE6)

def tint(c,p):
    return RGBColor(round(255*(1-p)+c[0]*p), round(255*(1-p)+c[1]*p), round(255*(1-p)+c[2]*p))
def shade_on(base,c,p):
    return RGBColor(round(base[0]*(1-p)+c[0]*p), round(base[1]*(1-p)+c[1]*p), round(base[2]*(1-p)+c[2]*p))

# strengthened tints (projector-safe)
TEAL_T   = tint(TEAL,0.07);  STEEL_T  = tint(STEEL,0.07)
TEAL_B   = tint(TEAL,0.40);  STEEL_B  = tint(STEEL,0.40)
GRAPH_B  = tint(GRAPHITE,0.25)
GRAPH_T  = tint(GRAPHITE,0.045)
STEEL_MID= tint(STEEL,0.38)                      # pre-step bars
TEAL_LT  = RGBColor(0x8F,0xC6,0xC6)              # pathway names on deep navy
STEEL_LT = RGBColor(0xA9,0xC6,0xE2)
NAVY_CARD_L = shade_on(DEEP_NAVY,WHITE,0.07)     # tonal card on dark field
NAVY_CARD_B = shade_on(DEEP_NAVY,WHITE,0.22)
TONAL_LN    = shade_on(DEEP_NAVY,WHITE,0.12)     # facet line on deep navy

HEAD="Aptos Display"; BODY="Aptos"; SERIF="Spectral Light"

SW,SH=Inches(13.333),Inches(7.5)

def new_pres():
    prs=Presentation(); prs.slide_width,prs.slide_height=SW,SH
    return prs
def add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

# ---------------- primitives ----------------
def rect(s,x,y,w,h,fill=None,line=None,line_w=None,shape=MSO_SHAPE.RECTANGLE,alpha=None):
    sp=s.shapes.add_shape(shape,x,y,w,h); sp.shadow.inherit=False
    if fill is None: sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb=fill
        if alpha is not None:
            clr=sp.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            a=clr.makeelement(qn('a:alpha'),{'val':str(int(alpha*100000))}); clr.append(a)
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=line_w or Pt(1)
    return sp

def card(s,x,y,w,h,fill=WHITE,line=None,line_w=Pt(0.75),radius=Inches(0.09),shadow=True):
    sp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h)
    sp.shadow.inherit=False
    try: sp.adjustments[0]=float(radius)/float(min(w,h))
    except Exception: pass
    sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=line_w
    if shadow:
        el=sp._element.spPr.find(qn('a:effectLst'))
        o=etree.SubElement(el,qn('a:outerShdw'),
            {'blurRad':'63500','dist':'25400','dir':'5400000','rotWithShape':'0'})
        c=etree.SubElement(o,qn('a:srgbClr'),{'val':'002147'})
        etree.SubElement(c,qn('a:alpha'),{'val':'16000'})
    return sp

def hrule(s,x,y,w,weight=0.75,color=MIST): return rect(s,x,y,w,Pt(weight),fill=color)

def set_run(r,text,size,font,color,bold=False,italic=False,track=None,caps=False):
    r.text=text; r.font.size=Pt(size); r.font.name=font
    r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color
    rPr=r._r.get_or_add_rPr()
    for tag in ("a:latin","a:cs"):
        e=rPr.find(qn(tag))
        if e is None: e=rPr.makeelement(qn(tag),{}); rPr.append(e)
        e.set("typeface",font)
    if track is not None: rPr.set("spc",str(int(track*100)))
    if caps: rPr.set("cap","all")

def textbox(s,x,y,w,h,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    return tb,tf
def para(tf,first=False):
    return tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
def line(tf,text,size,font=BODY,color=GRAPHITE,bold=False,italic=False,align=PP_ALIGN.LEFT,
         sb=0,sa=6,first=False,track=None,caps=False,ls=None):
    p=para(tf,first=first); p.alignment=align
    p.space_before=Pt(sb); p.space_after=Pt(sa)
    if ls: p.line_spacing=ls
    r=p.add_run(); set_run(r,text,size,font,color,bold=bold,italic=italic,track=track,caps=caps)
    return p

# ---------------- furniture & shared blocks ----------------
def furniture(s,page,dark=False):
    fg=WHITE_D if dark else SLATE
    tb,tf=textbox(s,Inches(0.7),Inches(7.06),Inches(5),Inches(0.3),anchor=MSO_ANCHOR.MIDDLE)
    line(tf,"Private and Confidential",8.5,color=fg,first=True,sa=0)
    tb,tf=textbox(s,Inches(11.6),Inches(7.06),Inches(1.03),Inches(0.3),anchor=MSO_ANCHOR.MIDDLE)
    line(tf,str(page),8.5,color=fg,align=PP_ALIGN.RIGHT,first=True,sa=0)
    s.shapes.add_picture(A+("/logo_white.png" if dark else "/logo_navy.png"),
                         Inches(11.90),Inches(0.36),height=Inches(0.38))

def kicker(s,x,y,text,color=GRAPHITE):
    tb,tf=textbox(s,x,y,Inches(10),Inches(0.28))
    line(tf,text,10,color=color,bold=True,first=True,track=1.6,caps=True,sa=0)

def title(s,x,y,text,size=24,color=NAVY,w=Inches(11.9)):
    tb,tf=textbox(s,x,y,w,Inches(1.0))
    line(tf,text,size,font=HEAD,color=color,first=True,ls=1.04,sa=0)
    return tb

def decision_kicker(s,x,y,text="For decision",dark=False):
    rect(s,x,y+Inches(0.05),Inches(0.11),Inches(0.11),fill=GOLD)
    tb,tf=textbox(s,x+Inches(0.24),y,Inches(9),Inches(0.28))
    line(tf,text,10,color=(WHITE if dark else NAVY),bold=True,first=True,track=1.6,caps=True,sa=0)

def conclusion_band(s,x,y,w,lead,rest,h=Inches(0.60)):
    card(s,x,y,w,h,fill=DEEP_NAVY,line=None)
    tb,tf=textbox(s,x+Inches(0.28),y,w-Inches(0.56),h,anchor=MSO_ANCHOR.MIDDLE)
    p=para(tf,first=True); p.space_after=Pt(0)
    r=p.add_run(); set_run(r,lead+"  ",13,HEAD,WHITE)
    r=p.add_run(); set_run(r,rest,13,BODY,WHITE_D)

def pathway_marker(s,x,y,h,acc):  # for dense tables: marker, not row tint
    rect(s,x,y,Pt(3),h,fill=acc)

def bars(s,x,base_y,w,heights,fill_lo,fill_hi=None,step_idx=None,maxh=1.18,bw=Inches(0.38)):
    n=len(heights); gp=(w-bw*n)/(n-1)
    for i,hh in enumerate(heights):
        bh=Inches(hh*maxh)
        f=fill_hi if (step_idx is not None and i>=step_idx and fill_hi) else fill_lo
        rect(s,x+i*(bw+gp),base_y-bh,bw,bh,fill=f)
    hrule(s,x,base_y,w,weight=1.0,color=GRAPH_B)

def cover_title(s,x,y,w,roman,italic,size=44,color=WHITE):
    tb,tf=textbox(s,x,y,w,Inches(1.05))
    p=para(tf,first=True); p.alignment=PP_ALIGN.LEFT; p.space_after=Pt(0)
    r=p.add_run(); set_run(r,roman,size,SERIF,color)
    r=p.add_run(); set_run(r,italic,size,SERIF,color,italic=True)

def cover_B(s,title_roman,title_italic,subtitle,meta,entity_line):
    """Default Board cover: navy institutional, vertical crop, tonal facet."""
    rect(s,0,0,SW,SH,fill=DEEP_NAVY)
    rect(s,Inches(-2.1),Inches(4.15),Inches(6.4),Inches(4.6),fill=None,line=TONAL_LN,
         line_w=Pt(1.0),shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    s.shapes.add_picture(A+"/coverB_strip.jpg",Inches(9.36),0,width=Inches(3.973),height=SH)
    s.shapes.add_picture(A+"/logo_white.png",Inches(1.0),Inches(0.75),height=Inches(0.58))
    hrule(s,Inches(1.0),Inches(2.86),Inches(0.85),weight=1.5,color=GOLD)
    tb,tf=textbox(s,Inches(1.0),Inches(3.04),Inches(8),Inches(0.32))
    line(tf,meta,11,color=SLATE,bold=True,first=True,track=1.8,caps=True,sa=0)
    cover_title(s,Inches(0.98),Inches(3.44),Inches(8.2),title_roman,title_italic,size=42)
    tb,tf=textbox(s,Inches(1.0),Inches(4.58),Inches(7.5),Inches(1.1))
    line(tf,subtitle,13.5,color=WHITE_D,first=True,ls=1.32,sa=0)
    tb,tf=textbox(s,Inches(1.0),Inches(6.95),Inches(8),Inches(0.32))
    line(tf,entity_line,9,color=WHITE_D,first=True,sa=0)

def cover_A(s,title_roman,title_italic,subtitle,meta,entity_line):
    """Investor/marketing cover: full-bleed cinematic + controlled navy overlay."""
    s.shapes.add_picture(A+"/coverA_full.jpg",0,0,width=SW,height=SH)
    rect(s,0,0,SW,SH,fill=DEEP_NAVY,alpha=0.66)
    s.shapes.add_picture(A+"/logo_white.png",Inches(1.0),Inches(0.75),height=Inches(0.58))
    hrule(s,Inches(1.0),Inches(2.98),Inches(0.85),weight=1.5,color=GOLD)
    tb,tf=textbox(s,Inches(1.0),Inches(3.16),Inches(10.5),Inches(0.32))
    line(tf,meta,11,color=WHITE_D,bold=True,first=True,track=1.8,caps=True,sa=0)
    cover_title(s,Inches(0.98),Inches(3.56),Inches(11.4),title_roman,title_italic,size=46)
    tb,tf=textbox(s,Inches(1.0),Inches(4.78),Inches(8.7),Inches(1.0))
    line(tf,subtitle,14.5,color=WHITE_D,first=True,ls=1.32,sa=0)
    tb,tf=textbox(s,Inches(1.0),Inches(6.95),Inches(11.3),Inches(0.32))
    line(tf,entity_line,9,color=WHITE_D,first=True,sa=0)

def divider_typographic(s,number,name,note=None):
    """Typography-led section divider (Board/IC default)."""
    rect(s,0,0,SW,SH,fill=DEEP_NAVY)
    s.shapes.add_picture(A+"/logo_white.png",Inches(11.63),Inches(0.40),height=Inches(0.38))
    tb,tf=textbox(s,Inches(1.0),Inches(2.55),Inches(2.4),Inches(1.5))
    line(tf,number,60,font=SERIF,color=WHITE_D,first=True,sa=0)
    rect(s,Inches(1.02),Inches(3.95),Pt(28),Pt(1.5),fill=GOLD)
    tb,tf=textbox(s,Inches(2.9),Inches(2.72),Inches(9.4),Inches(1.6))
    line(tf,name,30,font=SERIF,color=WHITE,first=True,ls=1.1,sa=0)
    if note:
        tb,tf=textbox(s,Inches(2.92),Inches(4.05),Inches(8.6),Inches(0.8))
        line(tf,note,12.5,color=WHITE_D,first=True,ls=1.3,sa=0)
    tb,tf=textbox(s,Inches(1.0),Inches(7.06),Inches(5),Inches(0.3))
    line(tf,"Private and Confidential",8.5,color=WHITE_D,first=True,sa=0)

def load_disclaimers(path="disclaimers.json"):
    return json.load(open(path))

def disclaimer_slide(s,jur_name,paragraphs,page):
    """Jurisdiction disclaimer — wording preserved verbatim; typography refined only."""
    rect(s,0,0,SW,SH,fill=WHITE)
    LX=Inches(0.7); RW=Inches(11.93)
    kicker(s,LX,Inches(0.42),"Important notice · "+jur_name)
    hrule(s,LX,Inches(0.74),RW,weight=1.0,color=NAVY)
    body=[p for p in paragraphs if p.strip() and p.strip()!="DISCLAIMER"]
    colw=(RW-Inches(0.4))/2
    boxes=[(LX,Inches(0.92),colw,Inches(5.95)),(LX+colw+Inches(0.4),Inches(0.92),colw,Inches(5.95))]
    half=(len(body)+1)//2
    for (bx,by,bw,bh),chunk in zip(boxes,[body[:half],body[half:]]):
        tb,tf=textbox(s,bx,by,bw,bh)
        first=True
        for ptext in chunk:
            line(tf,ptext,7.5,color=GRAPHITE,first=first,ls=1.12,sa=5)
            first=False
    furniture(s,page)
